from __future__ import annotations

import json
from pathlib import Path

import cv2
import numpy as np
import openpyxl

from beewings.core.profiles import get_profile
from beewings.core.schema import Landmark, WingAnnotation, annotation_path, save_annotation
from beewings.pipeline.project import CropProject


def _annotate(cdir, cp, profile):
    lms = [Landmark(id=i, x=float(i * 5 + 1), y=float(i * 3 + 1)) for i in profile.ids]
    ann = WingAnnotation(image=cp.name, image_size=(80, 40),
                         profile=profile.name, landmarks=lms)
    save_annotation(ann, annotation_path(cp, cdir, profile.methodology_id))


def _make_project_with_crops(tmp_path, methodologies=("Алпатов 12 точек",)):
    root = tmp_path / "proj"
    scan = tmp_path / "8.jpg"
    cv2.imwrite(str(scan), np.full((50, 50, 3), 255, np.uint8))
    proj = CropProject.create(root, scans_root=str(tmp_path), scan_paths=[scan])
    proj.settings.profile = "Алпатов 12 точек"
    entry = proj.scans[0]
    entry.cropped = True
    cdir = proj.crops_dir(entry)
    cdir.mkdir(parents=True)
    for i in range(2):
        cp = cdir / f"8_crop_{i}.jpg"
        cv2.imwrite(str(cp), np.full((40, 80, 3), 255, np.uint8))
        for prof_name in methodologies:
            _annotate(cdir, cp, get_profile(prof_name))
    proj.save()
    return proj


def test_export_per_scan_writes_suffixed_files(tmp_path):
    from beewings.pipeline.export import export_per_scan
    proj = _make_project_with_crops(tmp_path)
    stats = export_per_scan(proj)

    sdir = Path(proj.root) / "8"
    assert sdir.is_dir()
    assert (sdir / "8.jpg").exists()              # the scan itself
    assert (sdir / "crops").is_dir()
    assert len(list((sdir / "crops").glob("*.jpg"))) == 2
    # methodology-suffixed artifacts
    assert (sdir / "8_alpatov.tps").exists()
    assert (sdir / "8_alpatov.xlsx").exists()
    assert (sdir / "8_alpatov.json").exists()
    # no tofilski annotations -> no tofilski files
    assert not (sdir / "8_tofilski.json").exists()

    assert stats["n_scans"] == 1


def test_export_separates_two_methodologies(tmp_path):
    from beewings.pipeline.export import export_per_scan
    proj = _make_project_with_crops(
        tmp_path, methodologies=("Алпатов 12 точек", "Тофильский 19 точек"))
    export_per_scan(proj)
    sdir = Path(proj.root) / "8"
    for mid in ("alpatov", "tofilski"):
        assert (sdir / f"8_{mid}.tps").exists()
        assert (sdir / f"8_{mid}.xlsx").exists()
        assert (sdir / f"8_{mid}.json").exists()

    a = json.loads((sdir / "8_alpatov.json").read_text(encoding="utf-8"))
    t = json.loads((sdir / "8_tofilski.json").read_text(encoding="utf-8"))
    assert a["methodology"] == "alpatov"
    assert t["methodology"] == "tofilski"
    assert len(a["wings"][0]["landmarks"]) == 12
    assert len(t["wings"][0]["landmarks"]) == 19
    # classical indices only for Alpatov
    assert "index_means" in a
    assert t["index_means"] == {}


def test_export_writes_dw_png_for_tofilski(tmp_path):
    from beewings.pipeline.export import export_per_scan
    from beewings.core.io_dw import read_dw_png
    proj = _make_project_with_crops(
        tmp_path, methodologies=("Алпатов 12 точек", "Тофильский 19 точек"))
    export_per_scan(proj)
    sdir = Path(proj.root) / "8"
    dw_files = sorted((sdir / "dw").glob("*.dw.png"))
    assert len(dw_files) == 2                       # one per annotated wing
    assert len(read_dw_png(dw_files[0])) == 19      # round-trips to 19 points


def test_export_formats_selection(tmp_path):
    from beewings.pipeline.export import export_per_scan
    proj = _make_project_with_crops(
        tmp_path, methodologies=("Алпатов 12 точек", "Тофильский 19 точек"))
    export_per_scan(proj, formats=["json", "dw"])  # only JSON + .dw.png
    sdir = Path(proj.root) / "8"
    assert (sdir / "8_tofilski.json").exists()
    assert (sdir / "dw").is_dir() and list((sdir / "dw").glob("*.dw.png"))
    # unchecked formats are not written
    assert not (sdir / "8_tofilski.tps").exists()
    assert not (sdir / "8_alpatov.xlsx").exists()
    # scan + crops always copied regardless of format selection
    assert (sdir / "8.jpg").exists()
    assert (sdir / "crops").is_dir()


def test_export_no_dw_png_without_tofilski(tmp_path):
    from beewings.pipeline.export import export_per_scan
    proj = _make_project_with_crops(tmp_path)  # Alpatov only
    export_per_scan(proj)
    assert not (Path(proj.root) / "8" / "dw").exists()


def test_export_json_has_metadata_and_wings(tmp_path):
    from beewings.pipeline.export import export_per_scan
    proj = _make_project_with_crops(tmp_path)
    export_per_scan(proj)
    data = json.loads((Path(proj.root) / "8" / "8_alpatov.json").read_text(encoding="utf-8"))
    assert data["scan"] == "8"
    assert data["profile"] == "Алпатов 12 точек"
    assert data["methodology"] == "alpatov"
    assert data["n_wings"] == 2
    assert "exported_at" in data
    assert len(data["wings"]) == 2


def test_export_xlsx_has_rows(tmp_path):
    from beewings.pipeline.export import export_per_scan
    proj = _make_project_with_crops(tmp_path)
    export_per_scan(proj)
    wb = openpyxl.load_workbook(Path(proj.root) / "8" / "8_alpatov.xlsx")
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert "wing" in header and "x1" in header and "y1" in header
    assert ws.max_row == 3  # header + 2 wings


def test_export_skips_scans_without_annotations(tmp_path):
    from beewings.pipeline.export import export_per_scan
    proj = _make_project_with_crops(tmp_path)
    scan2 = tmp_path / "9.jpg"
    cv2.imwrite(str(scan2), np.full((50, 50, 3), 255, np.uint8))
    proj.sync_new_scans()
    stats = export_per_scan(proj)
    assert stats["n_scans"] == 1
    assert not (Path(proj.root) / "9").exists()
